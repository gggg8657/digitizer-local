#!/usr/bin/env python3
"""LLM 없이 검증: 참값을 아는 합성 그래프(선형 2계열·로그축·산점; Pillow + 있으면 matplotlib)를 그려
축·눈금 검출 → 눈금 값 짝짓기(가짜 VLM) → 보정 → 색 추출 → 오차(%)를 잰다. 숫자 파싱·짝짓기 견고성·CSV/XLSX·이력·HTTP·PDF.
WORKSPACE 는 임시 폴더로 바꿔 실데이터 폴더에 흔적을 남기지 않는다.
python3 selftest.py          # 정확도 표를 출력.  SELFTEST_OUT=폴더 → 합성 그림·겹쳐 보기 저장"""
import base64, io, json, math, os, re, shutil, sys, tempfile, threading, urllib.request  # noqa: E401

TMP = tempfile.mkdtemp(prefix="digitizer-selftest-")
os.environ["WORKSPACE"] = TMP
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np  # noqa: E402
import app, digitize as D, synth as S  # noqa: E401,E402

OUT = os.environ.get("SELFTEST_OUT")
ROWS = []


def fs_err(cal, d, f, log):
    """곡선 오차: 축 전체 범위(로그축은 자릿수) 대비 % — rms·max, 로그축은 상대오차 %도"""
    d = np.asarray(d)
    yt = f(d[:, 0])
    if log:
        span = math.log10(cal["y"]["v2"]) - math.log10(cal["y"]["v1"])
        e = np.abs(np.log10(d[:, 1]) - np.log10(yt)) / abs(span) * 100
        rel = np.abs(d[:, 1] / yt - 1) * 100
        return e, rel
    return np.abs(d[:, 1] - yt) / abs(cal["y"]["v2"] - cal["y"]["v1"]) * 100, None


def case(name, png, tr, funcs, mode="line", log=False, lim=(0.4, 1.5), vlm=True):
    a = D.load(png)
    if OUT:
        os.makedirs(OUT, exist_ok=True)
        open(os.path.join(OUT, name + ".png"), "wb").write(png)
    r = D.auto_calibrate(a, tr["xticks"], tr["yticks"]) if vlm else None
    assert r and r["calib"], f"{name}: 자동 보정 실패 {r and r.get('error')}"
    cal = r["calib"]
    assert cal["y"]["log"] == bool(tr["calib"]["y"]["log"]), f"{name}: 로그축 판단 틀림"
    # 자동 보정 자체 오차: 그림 영역 모서리를 참 보정·자동 보정으로 바꿔 비교 (축 범위 대비 %)
    bx = tr["box"]
    P = np.array([[bx["x0"], bx["y1"]], [bx["x1"], bx["y0"]]], float)
    ta, tb = D.to_data(tr["calib"], P[:, 0], P[:, 1])
    aa, ab = D.to_data(cal, P[:, 0], P[:, 1])
    fx = (lambda v: np.log10(v)) if tr["calib"]["x"]["log"] else (lambda v: v)
    fy = (lambda v: np.log10(v)) if tr["calib"]["y"]["log"] else (lambda v: v)
    cerr = max(np.max(np.abs(fx(aa) - fx(ta))) / abs(fx(ta[1]) - fx(ta[0])), np.max(np.abs(fy(ab) - fy(tb))) / abs(fy(tb[1]) - fy(tb[0]))) * 100
    assert cerr < 0.5, f"{name}: 자동 보정 오차 {cerr:.2f}%"
    for s, f in zip(tr["series"], funcs):
        col = s["color"] if isinstance(s["color"], str) else D.rgb2hex(s["color"])
        res = D.run_series(a, cal, r["box"], dict(color=col, tol=70, mode=mode, how="n", n=100))
        d = res["data"]
        if mode == "line":
            assert len(d) >= 90, f"{name}/{s['name']}: 점 {len(d)}개"
            e, rel = fs_err(tr["calib"], d, f, log)
            rms, mx = float(np.sqrt((e ** 2).mean())), float(e.max())
            assert rms < lim[0] and mx < lim[1], f"{name}/{s['name']}: 오차 rms {rms:.3f}% max {mx:.3f}%"
            ROWS.append((name, s["name"], len(d), cerr, rms, mx, None if rel is None else float(np.sqrt((rel ** 2).mean()))))
        else:
            tx, ty = np.asarray(s["x"], float), np.asarray(s["y"], float)
            d = np.asarray(d)
            assert len(d) == len(tx), f"{name}/{s['name']}: 점 {len(d)}개 (참 {len(tx)}개)"
            o = np.argsort(tx)
            gx = np.log10 if tr["calib"]["x"]["log"] else (lambda v: v)
            gy = np.log10 if tr["calib"]["y"]["log"] else (lambda v: v)
            ex = np.abs(gx(d[:, 0]) - gx(tx[o])) / abs(gx(tr["calib"]["x"]["v2"]) - gx(tr["calib"]["x"]["v1"])) * 100
            ey = np.abs(gy(d[:, 1]) - gy(ty[o])) / abs(gy(tr["calib"]["y"]["v2"]) - gy(tr["calib"]["y"]["v1"])) * 100
            mx = float(max(ex.max(), ey.max()))
            assert mx < lim[1], f"{name}/{s['name']}: 점 위치 오차 max {mx:.3f}%"
            ROWS.append((name, s["name"], len(d), cerr, float(np.sqrt(((ex ** 2 + ey ** 2) / 2).mean())), mx, None))
    return a, r


try:
    assert app.WS == TMP, "WORKSPACE 가 임시 폴더가 아님"
    xs = np.linspace(0, 10, 500)

    # 1) Pillow — 선형 2계열 / 로그 y / 산점 (어느 서버에서나)
    png, tr = S.pillow_plot([dict(name="sin", color=(31, 119, 180), kind="line", x=xs, y=np.sin(xs)),
                             dict(name="0.6cos", color=(214, 39, 40), kind="line", x=xs, y=0.6 * np.cos(xs))],
                            (0, 10), (-1.5, 1.5), [0, 2, 4, 6, 8, 10], [-1.5, -1, -0.5, 0, 0.5, 1, 1.5])
    a, r = case("pillow-선형", png, tr, [np.sin, lambda x: 0.6 * np.cos(x)])
    assert r["ticks"]["x_kind"] == "out" and len(r["ticks"]["x"]) == 6 and len(r["labels"]["x"]) == 6 and len(r["labels"]["y"]) == 7
    cols = [c["color"] for c in D.suggest_colors(a, r["box"])]
    assert cols[:2] == ["#1f77b4", "#d62728"], cols
    png, tr = S.pillow_plot([dict(name="exp", color=(44, 160, 44), kind="line", x=xs, y=np.exp(0.8 * xs))],
                            (0, 10), (0.1, 1e4), [0, 2, 4, 6, 8, 10], [0.1, 1, 10, 100, 1000, 10000], ylog=True)
    case("pillow-로그y", png, tr, [lambda x: np.exp(0.8 * x)], log=True)
    px_ = np.array([0.7, 1.9, 3.1, 4.4, 5.2, 6.8, 7.5, 9.1]); py_ = 2 + 0.25 * px_ ** 2
    png, tr = S.pillow_plot([dict(name="점", color=(148, 103, 189), kind="points", x=px_, y=py_)], (0, 10), (0, 30), [0, 2, 4, 6, 8, 10], [0, 10, 20, 30])
    case("pillow-산점", png, tr, [None], mode="points", lim=(0.3, 0.4))

    # 2) matplotlib (있으면) — 안티앨리어싱·격자·범례가 있는 '논문 같은' 그림, 100·200 dpi
    try:
        import matplotlib  # noqa: F401
        has_mpl = True
    except ImportError:
        has_mpl = False
    if has_mpl:
        for dpi in (100, 200):
            png, tr = S.mpl_plot([dict(name="sin", color="#1f77b4", kind="line", x=xs, y=np.sin(xs)),
                                  dict(name="0.5cos", color="#d62728", kind="line", x=xs, y=0.5 * np.cos(xs), ls="--")], dpi=dpi)
            case(f"mpl-선형-{dpi}dpi", png, tr, [np.sin, lambda x: 0.5 * np.cos(x)])
            png, tr = S.mpl_plot([dict(name="exp", color="#2ca02c", kind="line", x=xs, y=np.exp(0.8 * xs)),
                                  dict(name="3exp", color="#ff7f0e", kind="line", x=xs, y=3 * np.exp(0.5 * xs))], ylog=True, dpi=dpi)
            case(f"mpl-로그y-{dpi}dpi", png, tr, [lambda x: np.exp(0.8 * x), lambda x: 3 * np.exp(0.5 * x)], log=True)
            png, tr = S.mpl_plot([dict(name="A", color="#9467bd", kind="points", x=px_, y=py_),
                                  dict(name="B", color="#2ca02c", kind="points", marker="s", x=px_ + 0.3, y=28 - py_)], dpi=dpi,
                                 xlim=(0, 10), ylim=(0, 32))
            case(f"mpl-산점-{dpi}dpi", png, tr, [None, None], mode="points", lim=(0.3, 0.5))
        # 로그-로그 (x·y 모두 로그, 거듭제곱 법칙)
        xx = np.logspace(-1, 2, 300)
        png, tr = S.mpl_plot([dict(name="x^-1.5", color="#1f77b4", kind="line", x=xx, y=5 * xx ** -1.5)], xlog=True, ylog=True, dpi=150)
        a, r = case("mpl-로그로그-150dpi", png, tr, [lambda x: 5 * x ** -1.5], log=True)
        assert r["calib"]["x"]["log"] and r["calib"]["y"]["log"]
        # 논문형: 세리프·안쪽/작은 눈금·네 변 눈금·틀 없는 범례 — 선+표식 3계열(곡선·점 둘 다), 빈 원 + 점선 맞춤선
        png, tr = S.paper_plot("A")
        a, r = case("논문형A-로그y", png, tr, [s["f"] for s in tr["series"]], log=True, lim=(0.5, 2.0))
        assert r["ticks"]["y_kind"] == "in" and D.ticks_look_log(r["ticks"]["y"]) is True
        for s in tr["series"]:  # 선으로 이어진 표식 → 점 방식으로도 13개, 범례 견본은 빠짐
            pts = D.run_series(a, r["calib"], r["box"], dict(color=s["color"], tol=70, mode="points"))["data"]
            assert len(pts) == 13, (s["name"], len(pts))
            e = max(abs(math.log10(p[1] / q)) for p, q in zip(pts, s["y"])) / math.log10(30) * 100
            ROWS.append(("논문형A-표식", s["name"], len(pts), 0.0, float("nan"), e, None))
            assert e < 1.0, (s["name"], e)
        png, tr = S.paper_plot("B")
        case("논문형B-로그로그산점", png, tr, [None], mode="points", lim=(0.4, 0.5))
        # 흐린 확대본(저해상도 그림이 든 PDF 를 300dpi 로 뽑은 꼴): 0.9배 줄였다가 2.2배 키움 → 굵고 흐린 선·눈금, 점선에 갈린 빈 원
        from PIL import Image as _Im
        im = _Im.open(io.BytesIO(png))
        im = im.resize((int(im.width * 0.9), int(im.height * 0.9)), _Im.LANCZOS)
        sc = 0.9 * 2.2
        im = im.resize((int(im.width * 2.2), int(im.height * 2.2)), _Im.BICUBIC)
        bb = io.BytesIO(); im.convert("RGB").save(bb, "PNG")
        trb = dict(tr, box={k: v * sc for k, v in tr["box"].items()},
                   calib={k: dict(v, p1=[q * sc for q in v["p1"]], p2=[q * sc for q in v["p2"]]) for k, v in tr["calib"].items()})
        case("논문형B-흐린확대", bb.getvalue(), trb, [None], mode="points", lim=(0.6, 0.8))
        png, tr = S.paper_plot("E", 130)  # 엑셀형: 세로축선 없음 → 가로 격자 왼쪽 끝을 축으로
        a, r = case("엑셀형-격자만", png, tr, [s["f"] for s in tr["series"]], lim=(0.4, 1.5))
        assert r["axes"]["synth_y"] and r["ticks"]["y_kind"] == "grid"
        for s in tr["series"]:  # 축에 걸친 첫 표식·범례 견본 포함해도 17개
            pts = D.run_series(a, r["calib"], r["box"], dict(color=s["color"], tol=60, mode="points"))["data"]
            assert len(pts) == 17, (s["name"], len(pts))
            e = max(abs(p[1] - q) for p, q in zip(pts, s["y"])) / 1600 * 100
            ROWS.append(("엑셀형-표식", s["name"], len(pts), 0.0, float("nan"), e, None))
            assert e < 0.8, (s["name"], e)

    # 3) 짝짓기 견고성: VLM 이 눈금 하나를 빠뜨리거나 더 읽어도 / 로그 판단 / 숫자 표기
    pos = [100.5, 200.5, 300.5, 400.5, 500.5]
    c = D.pair_axis([2, 4, 6, 8], pos, pos)          # 첫 눈금(0)을 못 읽음 → 등간격이라 어느 칸인지 모름 → ambiguous
    assert c["n"] == 4 and c["ambiguous"] and not c["log"], c
    c = D.pair_axis([None, 2, 4, 6, 8], pos, pos, fixed=True)  # 덩어리별로 읽으면(1:1) 확정
    assert c["n"] == 4 and abs(c["p1"] - 200.5) < 1e-6 and c["v1"] == 2 and not c["ambiguous"], c
    c = D.pair_axis([0, 2, 4, 6, 8], pos, pos)
    assert c["n"] == 5 and c["resid_px"] < 1e-6 and not c["ambiguous"]
    c = D.pair_axis([0, 2, 4, 7, 8], pos, pos)        # 하나 잘못 읽음 → 잔차로 드러남
    assert c["resid_px"] > 1.5 or c["n"] < 5
    c = D.pair_axis([1, 10, 100, 1000, 10000], pos, pos)  # 로그 간격 → 로그축
    assert c["log"] and c["resid_px"] < 1e-6
    for s_, v in (("1,000", 1000), ("−0.5", -0.5), ("10^3", 1000), ("10^-2", 0.01), ("10⁻²", 0.01), ("1e-3", 0.001), ("2×10^3", 2000), (5, 5.0)):
        assert abs(app.num(s_) - v) < 1e-12, (s_, app.num(s_))

    # 4) 보정 수식: 왕복·기운 축(2°)·잘못된 입력
    cal = {"x": {"p1": [100, 400], "p2": [700, 400], "v1": 0, "v2": 60, "log": False},
           "y": {"p1": [100, 400], "p2": [100, 100], "v1": 1, "v2": 1000, "log": True}}
    X, Y = D.to_data(cal, [400, 250], [250, 325])
    assert np.allclose(X, [30, 15]) and np.allclose(Y, [10 ** 1.5, 10 ** 0.75])
    PX, PY = D.to_pixel(cal, X, Y)
    assert np.allclose(PX, [400, 250]) and np.allclose(PY, [250, 325])
    th = math.radians(2)
    rot = lambda x, y: [x * math.cos(th) - y * math.sin(th), x * math.sin(th) + y * math.cos(th)]  # noqa: E731
    calr = {"x": {"p1": rot(100, 400), "p2": rot(700, 400), "v1": 0, "v2": 60, "log": False},
            "y": {"p1": rot(100, 400), "p2": rot(100, 100), "v1": 0, "v2": 30, "log": False}}
    X, Y = D.to_data(calr, *zip(rot(400, 250)))
    assert abs(X[0] - 30) < 1e-9 and abs(Y[0] - 15) < 1e-9, (X, Y)
    for bad in ({"v1": 5, "v2": 5}, {"v1": -1, "v2": 10, "log": True}):
        try:
            D.matrix({"x": dict(cal["x"], **bad), "y": cal["y"]})
            raise AssertionError("잘못된 보정이 통과됨")
        except ValueError:
            pass

    # 5) 출력 간격: x 간격 0.5 → 0, 0.5, … 정확히 / 로그축 점 수는 로그 간격
    png, tr = S.pillow_plot([dict(name="lin", color=(31, 119, 180), kind="line", x=xs, y=0.2 * xs + 0.5)], (0, 10), (0, 3), [0, 2, 4, 6, 8, 10], [0, 1, 2, 3])
    a = D.load(png)
    res = D.run_series(a, tr["calib"], tr["box"], dict(color="#1f77b4", tol=60, how="step", step=0.5))
    xs_ = [p[0] for p in res["data"]]
    assert np.allclose(np.diff(xs_), 0.5) and all(abs(x / 0.5 - round(x / 0.5)) < 1e-9 for x in xs_) and len(xs_) >= 19, xs_[:4]
    assert max(abs(p[1] - (0.2 * p[0] + 0.5)) for p in res["data"]) < 0.02
    calL = dict(tr["calib"], x=dict(tr["calib"]["x"], v1=1, v2=1000, log=True))
    d = D.resample(calL, res["px"], "n", 7)
    assert len(d) == 7 and np.allclose(np.diff(np.log10([p[0] for p in d])), np.log10(d[1][0] / d[0][0]))
    # 범례·글자 조각 무시: 같은 색 짧은 선(범례 견본)을 그림 위쪽에 그려도 곡선만
    from PIL import Image, ImageDraw
    im = Image.open(io.BytesIO(png)).convert("RGB")
    ImageDraw.Draw(im).line([(110, 60), (150, 60)], fill=(31, 119, 180), width=2)
    b = io.BytesIO(); im.save(b, "PNG")
    res2 = D.run_series(D.load(b.getvalue()), tr["calib"], tr["box"], dict(color="#1f77b4", tol=60, how="all", skip_legend=False))
    assert max(abs(p[1] - (0.2 * p[0] + 0.5)) for p in res2["data"]) < 0.03, "범례 견본선이 곡선에 섞임"
    # 색이 없으면 경고
    res3 = D.run_series(a, tr["calib"], tr["box"], dict(color="#00ff00", tol=20))
    assert not res3["data"] and res3["warn"]

    # 6) 앱: 저장·가짜 VLM 자동 보정·추출·CSV/XLSX·이력
    png, tr = S.pillow_plot([dict(name="sin", color=(31, 119, 180), kind="line", x=xs, y=np.sin(xs)),
                             dict(name="cos", color=(214, 39, 40), kind="line", x=xs, y=np.cos(xs))],
                            (0, 10), (-1.5, 1.5), [0, 2, 4, 6, 8, 10], [-1.5, -1, -0.5, 0, 0.5, 1, 1.5])
    info = app.save_image(png)
    assert re.fullmatch(r"[0-9a-f]{16}", info["id"]) and info["w"] == 800 and app.save_image(png)["id"] == info["id"]
    seen = []

    def fake_vlm(prompt, images, model=None):
        seen.append(len(images))
        if "numbered row" in prompt:
            return json.dumps({"labels": {}})
        return json.dumps({"x_ticks": ["0", "2", "4", "6", "8", "10"], "y_ticks": ["−1.5", "−1", "−0.5", "0", "0.5", "1", "1.5"],
                           "x_scale": "linear", "y_scale": "linear", "x_label": "시간 (s)", "y_label": "변위 (mm)",
                           "legend": [{"name": "cos", "color": "red"}, {"name": "sin", "color": "blue"}]})
    app.chat_vision = fake_vlm
    r = app.auto(info["id"])
    d = r["detect"]
    assert r["error"] is None and seen == [3, 1, 1], (r["error"], seen)  # 전체+x 글자 띠+y 글자 띠, 그다음 덩어리별 x·y
    assert d["calib"] and abs(d["calib"]["y"]["v1"] + 1.5) < 1e-9 and d["fit"]["x"]["resid_px"] < 1
    assert [(s["name"], s["color"]) for s in d["series"]] == [("cos", "#d62728"), ("sin", "#1f77b4")], d["series"]
    # VLM 이 x 눈금 하나(0)를 빠뜨림 → 덩어리별 다시 읽기(1:1)로 확정
    calls = []

    def fake_vlm2(prompt, images, model=None):
        calls.append(prompt[:20])
        if "numbered row" in prompt and len(calls) == 2:
            return json.dumps({"labels": {"1": 0, "2": 2, "3": "4", "4": 6, "5": 8, "6": 10}})
        if "numbered row" in prompt:
            return json.dumps({"labels": {"1": "−1.5", "2": None, "3": -0.5, "4": 0, "5": 0.5, "6": 1, "7": 1.5}})
        return json.dumps({"x_ticks": [2, 4, 6, 8, 10], "y_ticks": [-1.5, -1, -0.5, 0, 0.5, 1, 1.5], "legend": []})
    app.chat_vision = fake_vlm2
    r = app.auto(info["id"])
    assert len(calls) == 3 and r["detect"]["calib"]["x"]["v1"] == 0 and r["detect"]["fit"]["x"]["source"] == "each" and r["vlm"]["x_ticks_each"] == [0, 2, 4, 6, 8, 10], (calls, r["detect"]["calib"]["x"])
    assert abs(r["detect"]["calib"]["x"]["p1"][0] - d["calib"]["x"]["p1"][0]) < 0.6
    app.chat_vision = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("연결 실패"))
    r = app.auto(info["id"])
    assert "VLM 읽기 실패" in r["error"] and r["detect"]["calib"] is None and r["detect"]["ticks"]["x"]
    r = app.auto(info["id"], use_vlm=False, xvals=[0, 2, 4, 6, 8, 10], yvals=[-1.5, -1, -0.5, 0, 0.5, 1, 1.5])
    calib = r["detect"]["calib"]
    assert calib
    ex = app.extract(info["id"], calib, r["detect"]["box"], {"color": "#1f77b4", "tol": 60, "how": "n", "n": 21})
    assert len(ex["data"]) == 21 and ex["mask"].startswith("data:image/png;base64,")
    # 화면 흐름과 같게: 데이터 → 픽셀로 되돌려 저장 → 내보낼 때 다시 데이터
    PX, PY = D.to_pixel(calib, [p[0] for p in ex["data"]], [p[1] for p in ex["data"]])
    ser = [{"name": "sin", "px": list(zip(PX.tolist(), PY.tolist()))}, {"name": "수동", "px": [[400.5, 260.5], [200.5, 100.5]]}]
    body, name, ctype = app.export(calib, ser, "csv", {"name": "그림 3"})
    txt = body.decode("utf-8-sig")
    rows = list(__import__("csv").reader(io.StringIO(txt)))
    assert name == "그림_3.csv" and rows[0] == ["sin x", "sin y", "수동 x", "수동 y"] and len(rows) == 22 and rows[3][2] == ""
    assert float(rows[1][2]) < float(rows[2][2]), "수동 점이 x 순으로 정렬되지 않음"
    assert all(abs(float(a) - float(b)) < 1e-5 * (1 + abs(b)) for a, b in zip([r_[0] for r_ in rows[1:]], [p[0] for p in ex["data"]]))
    body, name, ctype = app.export(calib, ser, "xlsx", {"name": "g"})
    if app._has("openpyxl"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(body))
        assert name == "g.xlsx" and wb.sheetnames == ["데이터", "sin", "수동", "보정"] and wb["데이터"].max_row == 22
    else:
        assert name == "g.csv"
    pid = app.save_project({"name": "그림 3", "image": info, "calib": calib, "box": r["detect"]["box"], "series": ser})
    h = app.history()
    assert h[0]["pid"] == pid and h[0]["series"][0]["n"] == 21 and app.load_project(pid)["name"] == "그림 3"
    assert app.save_project({"pid": pid, "name": "고침", "image": info, "calib": calib, "series": ser}) == pid and len(app.history()) == 1
    for bad in ("../x", "2026-01-01-zzzzzz"):
        try:
            app.load_project(bad)
            raise AssertionError("잘못된 id 통과")
        except (ValueError, FileNotFoundError):
            pass

    # 7) HTTP: 화면·메타·올리기·자동(값 직접)·추출·내보내기·PDF
    app.H.log_message = lambda *a: None
    srv = app.ThreadingHTTPServer(("127.0.0.1", 0), app.H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}/"

    def post(p, body, raw=False):
        rq = urllib.request.Request(base + p, json.dumps(body).encode(), {"Content-Type": "application/json"})
        with urllib.request.urlopen(rq) as f:
            return (f.read(), dict(f.headers)) if raw else json.load(f)
    html = urllib.request.urlopen(base).read().decode()
    assert "그래프 디지타이저" in html and "data-sig" in html
    assert "xlsx" in json.load(urllib.request.urlopen(base + "api/meta"))
    du = "data:image/png;base64," + base64.b64encode(png).decode()
    info2 = post("api/image", {"data": du})
    assert info2["id"] == info["id"]
    assert urllib.request.urlopen(base + f"api/img/{info['id']}.png").read()[:4] == b"\x89PNG"
    r = post("api/auto", {"id": info["id"], "vlm": False, "xvals": ["0", "10"], "yvals": ["-1.5", "1.5"]})
    assert r["detect"]["calib"], r
    ex = post("api/extract", {"id": info["id"], "calib": r["detect"]["calib"], "box": r["detect"]["box"], "series": {"color": "#d62728", "tol": 60, "mode": "line", "how": "all"}})
    assert len(ex["data"]) > 300
    b, hd = post("api/export", {"calib": r["detect"]["calib"], "format": "csv", "series": [{"name": "c", "px": ex["px"][:5]}]}, raw=True)
    assert b.startswith("﻿c x,c y".encode()) and "attachment" in hd["Content-Disposition"] and hd.get("X-Author")
    if shutil.which("pdftoppm"):
        from PIL import Image as _I
        pb = io.BytesIO()
        _I.open(io.BytesIO(png)).convert("RGB").save(pb, "PDF", resolution=100)
        pdf = post("api/pdf", {"data": "data:application/pdf;base64," + base64.b64encode(pb.getvalue()).decode()})
        assert pdf["pages"] == 1
        pg = post("api/pdfpage", {"id": pdf["id"], "page": 1, "dpi": 200})
        assert pg["w"] == 1600, pg
        assert urllib.request.urlopen(base + f"api/pdfpage/{pdf['id']}/1.png?dpi=24").read()[:4] == b"\x89PNG"
    try:
        post("api/auto", {"id": "../../etc"})
        raise AssertionError("잘못된 id 통과")
    except urllib.error.HTTPError as e:
        assert e.code == 500
    srv.shutdown()

    # 9) 막대그래프: 값축만 보정 → 막대·쌓인 조각·무늬·오차 막대·가로·로그 (참값 대비, 값축 범위의 %)
    if app._has("matplotlib"):
        for kind in ("simple", "stacked", "hatched", "error", "hbar", "log"):
            for dpi in (80, 150):
                png, tr = S.bar_plot(kind, dpi)
                a = D.load(png)
                if OUT:
                    open(os.path.join(OUT, f"bar_{kind}_{dpi}.png"), "wb").write(png)
                det = D.auto_calibrate(a, None, None)
                xv, yv = (tr["ticks"], None) if tr["orient"] == "h" else (None, tr["ticks"])
                br = app.bar_auto(a, det, None, {}, xv, yv, use_vlm=False)
                assert br["ok"] and br["orient"] == tr["orient"], (kind, dpi, br.get("error"), br.get("orient"))
                truth = tr["bars"][::-1] if tr["orient"] == "h" else tr["bars"]  # 가로 막대는 위 → 아래
                assert len(br["bars"]) == len(truth), f"{kind}/{dpi}: 막대 {len(br['bars'])}개 (참 {len(truth)}개)"
                g = math.log10 if tr["log"] else (lambda v: v)
                span = abs(g(tr["lim"][1]) - g(tr["lim"][0]))
                errs, smap = [], {}
                for b, t in zip(br["bars"], truth):
                    tops, err = app.bar_values(br["calib"], b, br["orient"])
                    assert len(tops) == len(t["tops"]), f"{kind}/{dpi}: 조각 {[round(x, 2) for x in tops]} (참 {t['tops']})"
                    errs += [abs(g(x) - g(y)) / span * 100 for x, y in zip(tops, t["tops"])]
                    for gt, got in zip(t["series"], b["series"]):  # 같은 참 계열 → 같은 검출 계열 (1:1)
                        assert smap.setdefault(gt, got) == got, f"{kind}/{dpi}: 계열 섞임"
                    assert bool(t["err"]) == bool(err), f"{kind}/{dpi}: 오차 막대 {err} (참 {t['err']})"
                    if err:
                        errs += [abs(err[0] - t["err"][1]) / span * 100, abs(err[1] - t["err"][0]) / span * 100]
                assert len(set(smap.values())) == len(smap) == tr["n_series"], f"{kind}/{dpi}: 계열 {smap}"
                mx = max(errs)
                assert mx <= 2.0, f"{kind}/{dpi}: 막대 값 오차 {mx:.2f}%"
                cerr = abs(br["fit"]["y" if tr["orient"] == "v" else "x"]["resid_px"])
                ROWS.append((f"막대-{kind}-{dpi}dpi", f"{len(smap)}계열", sum(len(b["ends"]) for b in br["bars"]), cerr,
                             float(np.sqrt(np.mean(np.square(errs)))), mx, None))
        # 범례 견본 ↔ 조각 (무늬·색): 쌓인 무늬 막대의 세 계열이 모두 범례에 짝지어짐
        png, tr = S.bar_plot("hatched", 150)
        a = D.load(png)
        br = app.bar_auto(a, D.auto_calibrate(a, None, None), None, {}, None, tr["ticks"], use_vlm=False)
        assert sum(s["legend"] for s in br["series"]) == 3, [(s["color"], s["legend"]) for s in br["series"]]
        # 앱 흐름: 가짜 VLM 이 '막대'라 하고 x 눈금은 비움 → 막대 모드 + 범주·범례 이름 + 표·CSV·XLSX
        png, tr = S.bar_plot("stacked", 110)
        info = app.save_image(png)

        def fake_bar(prompt, images, model=None):
            if "magenta tag" in prompt:
                return json.dumps({"bars": {"1": {"category": "Conv.", "group": "EV"}, "2": {"category": "Cir.", "group": "EV"},
                                            "3": {"category": "Conv.", "group": "Phone"}, "4": {"category": "Cir.", "group": "Phone"}}})
            if "chart legend" in prompt:
                return json.dumps({"labels": {"1": "S1", "2": "S2", "3": "S3"}})
            if "numbered row" in prompt:
                return json.dumps({"labels": {}})
            return json.dumps({"x_ticks": [], "y_ticks": [0, 20, 40, 60, 80, 100], "chart_type": "bar", "legend": []})
        app.chat_vision = fake_bar
        r = app.auto(info["id"])
        assert r["mode"] == "bar" and r["bar"]["ok"] and "막대 4개" in r["message"], r["message"]
        bb = r["bar"]
        assert [b["cat"] for b in bb["bars"]] == ["EV Conv.", "EV Cir.", "Phone Conv.", "Phone Cir."], [b["cat"] for b in bb["bars"]]
        assert sorted(s["name"] for s in bb["series"]) == ["S1", "S2", "S3"], bb["series"]
        head, rows, long_ = app.bar_table(bb["calib"], bb["bars"], bb["series"], bb["orient"])
        assert head[0] == "범주" and head[-3:] == ["막대 끝(합계)", "오차 +", "오차 −"] and len(rows) == 4 and len(long_) == 12
        assert abs(rows[0][-3] - 77) < 1 and abs(sum(rows[0][1:4]) - rows[0][-3]) < 1e-6, rows[0]
        body, name, _ = app.export(bb["calib"], bb["series"], "csv", {"name": "막대"}, bb["bars"], bb["orient"])
        assert name == "막대.csv" and body.decode("utf-8-sig").splitlines()[1].startswith("EV Conv.,")
        if app._has("openpyxl"):
            import openpyxl
            body, name, _ = app.export(bb["calib"], bb["series"], "xlsx", {"name": "막대"}, bb["bars"], bb["orient"])
            wb = openpyxl.load_workbook(io.BytesIO(body))
            assert wb.sheetnames == ["데이터", "막대 조각", "보정"] and wb["막대 조각"].max_row == 13
        # VLM 이 차트 종류를 말하지 않아도: x 눈금이 없고 막대가 있으면 막대 모드로 전환하고 이유를 알림
        app.chat_vision = lambda p, i, m=None: (json.dumps({"labels": {}}) if "numbered row" in p or "chart legend" in p else
                                                json.dumps({"bars": {}}) if "magenta tag" in p else
                                                json.dumps({"x_ticks": [], "y_ticks": [0, 20, 40, 60, 80, 100], "legend": []}))
        r = app.auto(info["id"])
        assert r["mode"] == "bar" and "범주형" in r["message"], r["message"]
        # 값축 숫자도 못 읽으면 조용히 끝나지 않고 이유를 말한다
        app.chat_vision = lambda p, i, m=None: json.dumps({"labels": {}} if "numbered row" in p else {"x_ticks": [], "y_ticks": [], "legend": []})
        r = app.auto(info["id"])
        assert r["message"] and not r["detect"].get("calib"), r["message"]
        # 여러 그래프가 있는 쪽: 그래프 둘 + 지도(빽빽한 그림) → 그래프 2개를 찾고 지도는 'image' 로 표시, 고른 뒤 자르면 원본 기록
        from PIL import Image
        p1, _ = S.bar_plot("stacked", 90)
        p2, _ = S.bar_plot("simple", 90)
        i1, i2 = Image.open(io.BytesIO(p1)).convert("RGB"), Image.open(io.BytesIO(p2)).convert("RGB")
        page = Image.new("RGB", (i1.width + i2.width + 40, i1.height + 360), "white")
        page.paste(i1, (0, 0))
        page.paste(i2, (i1.width + 40, 0))
        rng = np.random.default_rng(1)
        noise = (rng.random((300, 400, 3)) * 160 + 60).astype(np.uint8)
        map_ = Image.fromarray(noise)
        page.paste(map_, (120, i1.height + 30))
        from PIL import ImageDraw
        ImageDraw.Draw(page).rectangle([120, i1.height + 30, 520, i1.height + 330], outline="black", width=2)
        bp = io.BytesIO(); page.save(bp, "PNG")
        pinfo = app.save_image(bp.getvalue())
        app.chat_vision = fake_bar
        r = app.auto(pinfo["id"])
        kinds = sorted(p["kind"] for p in r["panels"])
        assert r["need_panel"] and kinds.count("plot") == 2 and "그래프 2개" in r["message"], (kinds, r["message"])
        pl = [p for p in r["panels"] if p["kind"] == "plot"]
        c = app.crop_image(pinfo["id"], pl[0]["crop"])
        assert c["parent"] == pinfo["id"] and app.image_meta(c["id"])["box"][0] == round(pl[0]["crop"][0])
        r = app.auto(c["id"])
        assert not r["need_panel"] and r["mode"] == "bar" and len(r["bar"]["bars"]) == 4, r["message"]

    # 8) 폐쇄망: 외부 CDN 없음, 저작권 표기
    ui = app.read(os.path.join(app.ROOT, "ui.html"))
    assert not re.search(r'<(script|link|img)[^>]+(src|href)="https?://', ui), "외부 CDN 참조 있음"
    _src = open(os.path.join(app.ROOT, "app.py"), encoding="utf-8").read()
    assert "wqkgMjAyNiBnZ2dnODY1NyDCtyBkb25nanVraW0uZGV2QGdtYWlsLmNvbQ==" in _src and "signed(" in _src and "X-Author" in _src, "저작권 표기 누락"
finally:
    shutil.rmtree(TMP, ignore_errors=True)

print(f"{'그림':22s} {'계열':8s} {'점':>4s} {'보정%':>6s} {'rms%':>6s} {'max%':>6s} {'상대rms%':>8s}")
for n, s, k, c, rms, mx, rel in ROWS:
    print(f"{n:22s} {s:8s} {k:4d} {c:6.3f} {'' if rms != rms else f'{rms:6.3f}':>6s} {mx:6.3f} {'' if rel is None else f'{rel:8.2f}'}")
print("(보정% = 자동 보정의 축 범위 대비 오차, rms/max% = 추출값의 축 범위(로그축은 자릿수) 대비 오차, 산점은 x·y 위치)")
print("(막대: 점 = 조각 수, 보정% 자리 = 값축 눈금 잔차(px), rms/max% = 막대·조각 끝·오차 끝 값의 값축 범위 대비 오차)")
print("selftest OK")
