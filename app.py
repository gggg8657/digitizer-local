#!/usr/bin/env python3
"""digitizer local — 그래프 그림 → 데이터 (WebPlotDigitizer 류 + AI 보조). CDN 없음.

좌표는 픽셀 처리 + 축 보정으로 결정론적으로 계산한다(digitize.py). 로컬 VLM 은 눈금 숫자·축 이름·범례 읽기만 돕는다.
  python3 app.py                       # http://localhost:8783
  python3 app.py --cli 그림.png 보정.json 계열.json   # 서버 없이 추출 → CSV (stdout)
LLM(선택, 이미지 입력 가능 모델): LLM_API=ollama|openai, LLM_BASE_URL, LLM_MODEL (또는 VISION_MODEL)
"""
import base64
import csv
import datetime
import hashlib
import io
import json
import os
import re
import secrets
import shutil
import subprocess
import sys
import tempfile
import threading
import urllib.error
import urllib.parse
import urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import numpy as np
from PIL import Image

import bars as B
import digitize as D

ROOT = os.path.dirname(os.path.abspath(__file__))
WS = os.environ.get("WORKSPACE") or os.path.join(ROOT, "_workspace")  # 포털이 AGENT_DATA/<도구> 로 모아 줌
LLM_API = os.environ.get("LLM_API", "ollama")            # ollama | openai (vLLM·LM Studio·llama.cpp 등)
LLM_BASE = os.environ.get("LLM_BASE_URL", "http://localhost:8000/v1" if LLM_API == "openai" else "http://localhost:11434").rstrip("/")
MODEL = os.environ.get("VISION_MODEL") or os.environ.get("LLM_MODEL", "gemma3:27b")
LLM_KEY = os.environ.get("LLM_API_KEY", "")
PORT = int(os.environ.get("PORT", "8783"))
MAX_UPLOAD = 40 * 1024 * 1024
LOCK = threading.Lock()
_CACHE = {}  # 이미지 id → numpy 배열 (최근 몇 개)


def read(p):
    with open(p, encoding="utf-8") as f:
        return f.read()


def ws(*p):
    d = os.path.join(WS, *p[:-1]) if len(p) > 1 else WS
    os.makedirs(d, exist_ok=True)
    return os.path.join(WS, *p)


# ── 이미지·PDF ─────────────────────────────────────────────────────────
def save_image(raw):
    if len(raw) > MAX_UPLOAD:
        raise ValueError("파일이 너무 큽니다 (40MB 이하)")
    im = Image.open(io.BytesIO(raw))
    im.load()
    if max(im.size) > 6000:
        im.thumbnail((6000, 6000))
    a = D.load(_png(im))
    b = _png(Image.fromarray(a))
    iid = hashlib.sha1(b).hexdigest()[:16]
    p = ws("images", iid + ".png")
    if not os.path.exists(p):
        with open(p, "wb") as f:
            f.write(b)
    return {"id": iid, "w": a.shape[1], "h": a.shape[0]}


def _png(im):
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def image(iid):
    if not re.fullmatch(r"[0-9a-f]{16}", iid or ""):
        raise ValueError("이미지 id 가 잘못되었습니다")
    if iid not in _CACHE:
        p = os.path.join(WS, "images", iid + ".png")
        if not os.path.exists(p):
            raise FileNotFoundError("이미지가 없습니다 — 다시 올려 주세요")
        if len(_CACHE) > 8:
            _CACHE.pop(next(iter(_CACHE)))
        _CACHE[iid] = D.load(p)
    return _CACHE[iid]


def save_pdf(raw):
    if raw[:5] != b"%PDF-":
        raise ValueError("PDF 파일이 아닙니다")
    if len(raw) > MAX_UPLOAD:
        raise ValueError("파일이 너무 큽니다 (40MB 이하)")
    pid = hashlib.sha1(raw).hexdigest()[:16]
    p = ws("pdf", pid + ".pdf")
    with open(p, "wb") as f:
        f.write(raw)
    return {"id": pid, "pages": pdf_pages(p)}


def pdf_pages(p):
    if shutil.which("pdfinfo"):
        out = subprocess.run(["pdfinfo", p], capture_output=True, text=True, timeout=60).stdout
        m = re.search(r"^Pages:\s+(\d+)", out, re.M)
        if m:
            return int(m.group(1))
    try:
        import fitz  # PyMuPDF
        with fitz.open(p) as d:
            return d.page_count
    except ImportError:
        raise RuntimeError("PDF 를 그림으로 바꿀 도구가 없습니다 (poppler-utils 의 pdftoppm 또는 pip PyMuPDF)")


def pdf_render(pid, page, dpi):
    if not re.fullmatch(r"[0-9a-f]{16}", pid or ""):
        raise ValueError("PDF id 가 잘못되었습니다")
    p = os.path.join(WS, "pdf", pid + ".pdf")
    if not os.path.exists(p):
        raise FileNotFoundError("PDF 가 없습니다")
    page, dpi = int(page), max(20, min(600, int(dpi)))
    if shutil.which("pdftoppm"):
        with tempfile.TemporaryDirectory() as td:
            subprocess.run(["pdftoppm", "-r", str(dpi), "-f", str(page), "-l", str(page), "-png", "-singlefile", p,
                            os.path.join(td, "o")], check=True, capture_output=True, timeout=180)
            with open(os.path.join(td, "o.png"), "rb") as f:
                return f.read()
    import fitz
    with fitz.open(p) as d:
        return d[page - 1].get_pixmap(dpi=dpi).tobytes("png")


# ── VLM (선택: 눈금 숫자·축 이름·범례 읽기) ─────────────────────────────────
VLM_PROMPT = """You read numbers printed on a scientific chart. Images: (1) the whole chart, (2) the strip under the x-axis
(x tick labels), (3) the strip left of the y-axis (y tick labels). Some may be missing.
Return ONLY a JSON object:
{"x_ticks": [numbers printed under the x-axis, left to right],
 "y_ticks": [numbers printed beside the LEFT y-axis, bottom to top],
 "y2_ticks": [numbers printed beside a second y-axis on the RIGHT side, bottom to top; [] if there is none],
 "x_scale": "linear" or "log", "y_scale": "linear" or "log",
 "x_label": "x-axis title with unit", "y_label": "y-axis title with unit", "title": "",
 "legend": [{"name": "series name exactly as printed", "color": "color word, e.g. red/blue/black", "style": "line|dashed|markers",
             "marker": "circle|square|triangle|diamond|none", "axis": "left" or "right" (which y-axis the series uses)}],
 "chart_type": "line" or "scatter" or "bar" (vertical bars) or "hbar" (horizontal bars) or "other"}
For bar charts the category axis has words, not numbers: give an empty list for that axis.
Rules: copy tick numbers exactly as printed (10^3 or 1e3 -> 1000, 10^-2 -> 0.01, "−0.5" -> -0.5, "1,000" -> 1000).
Only numbers that are actually printed; do not invent or extrapolate. Scale is "log" when labels grow by a constant
factor (1, 10, 100 or 10^-1, 10^0, 10^1). Empty list if there is no legend."""


def chat_vision(prompt, images, model=None):
    model = model or MODEL
    b64 = [base64.b64encode(b).decode() for b in images]
    try:
        if LLM_API == "openai":
            content = [{"type": "text", "text": prompt}] + [{"type": "image_url", "image_url": {"url": "data:image/png;base64," + x}} for x in b64]
            hdr = {"Content-Type": "application/json", **({"Authorization": f"Bearer {LLM_KEY}"} if LLM_KEY else {})}
            body = {"model": model, "temperature": 0, "messages": [{"role": "user", "content": content}]}
            req = urllib.request.Request(LLM_BASE + "/chat/completions", json.dumps(body).encode(), hdr)
            with urllib.request.urlopen(req, timeout=600) as r:
                return json.load(r)["choices"][0]["message"]["content"]
        body = {"model": model, "stream": False, "think": False, "format": "json", "options": {"temperature": 0},
                "messages": [{"role": "user", "content": prompt, "images": b64}]}
        for attempt in (0, 1):
            try:
                req = urllib.request.Request(LLM_BASE + "/api/chat", json.dumps(body).encode(), {"Content-Type": "application/json"})
                with urllib.request.urlopen(req, timeout=600) as r:
                    return json.load(r)["message"]["content"]
            except urllib.error.HTTPError as e:
                msg = e.read().decode(errors="replace")
                if attempt == 0 and "think" in msg:
                    body.pop("think")
                    continue
                raise RuntimeError(f"LLM HTTP {e.code}: {msg[:200]}")
    except urllib.error.URLError as e:
        raise RuntimeError(f"LLM 서버 연결 실패 ({LLM_BASE}): {e.reason}")


def parse_json(text):
    text = re.sub(r"<think>.*?</think>", "", text or "", flags=re.S)
    m = re.search(r"\{.*\}", text, re.S)
    if not m:
        raise ValueError("VLM 응답에 JSON 이 없습니다")
    return json.loads(m.group())


def num(v):
    """'1,000' '−0.5' '10^3' '1e-2' '10⁻²' → float"""
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("−", "-").replace("–", "-").replace(" ", "")
    sup = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁻", "0123456789-")
    m = re.fullmatch(r"(-?\d*\.?\d*)[×x*·]?10(?:(?:\^|\*\*)\{?(-?\d+(?:\.\d+)?)\}?|([⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+))", s)
    if m:
        mant = {"": 1.0, "-": -1.0}.get(m.group(1))
        mant = float(m.group(1)) if mant is None else mant
        return mant * 10 ** float(m.group(2) or m.group(3).translate(sup))
    return float(s)


def _crop_png(a, box, scale=1.0, maxw=1400):
    if not box:
        return None
    x0, y0, x1, y1 = [int(round(v)) for v in box]
    x0, y0, x1, y1 = max(0, x0 - 4), max(0, y0 - 4), min(a.shape[1], x1 + 5), min(a.shape[0], y1 + 5)
    if x1 - x0 < 3 or y1 - y0 < 3:
        return None
    im = Image.fromarray(a[y0:y1, x0:x1])
    s = min(scale, maxw / max(im.size))
    if s != 1:
        im = im.resize((max(1, int(im.width * s)), max(1, int(im.height * s))), Image.LANCZOS)
    return _png(im)


def vlm_read(a, det=None, model=None):
    """그림 전체 + 눈금 글자 띠(확대)를 VLM 에 보여 숫자·이름을 읽게 한다. 좌표는 묻지 않는다."""
    im = Image.fromarray(a)
    if max(im.size) > 1400:
        im.thumbnail((1400, 1400))
    imgs = [_png(im)]
    if det and det.get("labels"):
        lb = det["labels"]
        for band in (lb.get("x_band"), lb.get("y_band")):
            if band:
                h = band[3] - band[1] + 1
                c = _crop_png(a, band, scale=max(1.0, min(4.0, 48.0 / max(h, 1))) if band is lb.get("x_band") else
                              max(1.0, min(4.0, 160.0 / max(band[2] - band[0] + 1, 1))))
                if c:
                    imgs.append(c)
    raw = chat_vision(VLM_PROMPT, imgs, model)
    j = parse_json(raw)
    out = {"raw": raw, "x_label": j.get("x_label") or "", "y_label": j.get("y_label") or "", "title": j.get("title") or "",
           "x_log": str(j.get("x_scale", "")).lower().startswith("log"), "y_log": str(j.get("y_scale", "")).lower().startswith("log"),
           "legend": [x for x in (j.get("legend") or []) if isinstance(x, dict)],
           "chart_type": str(j.get("chart_type") or "").lower()}
    for k in ("x", "y", "y2"):
        vals = []
        for v in j.get(k + "_ticks") or []:
            try:
                vals.append(num(v))
            except (ValueError, TypeError):
                pass
        out[k + "_ticks"] = vals
    return out


BLOB_PROMPT = """Each numbered row [1], [2], ... of this image shows ONE tick label cut from a chart axis.
Read the number in each row exactly as printed (10^3 -> 1000, 10^-2 -> 0.01, "−0.5" -> -0.5).
Return ONLY JSON: {"labels": {"1": number or null, "2": number or null, ...}} with one entry per row; null if unreadable."""


def blob_sheet(a, boxes):
    """글자 덩어리를 하나씩 잘라 번호를 붙여 한 장으로 — VLM 이 위치별로 읽게 (개수 어긋남 방지)"""
    from PIL import ImageDraw
    crops = []
    for b in boxes:
        x0, y0, x1, y1 = [int(round(v)) for v in b]
        im = Image.fromarray(a[max(0, y0 - 3):y1 + 3, max(0, x0 - 3):x1 + 3])
        s = 44.0 / max(im.height, 1)
        crops.append(im.resize((max(1, int(im.width * s)), 44), Image.LANCZOS))
    W = 90 + max(c.width for c in crops) + 20
    sheet = Image.new("RGB", (W, 60 * len(crops) + 10), "white")
    d = ImageDraw.Draw(sheet)
    fn = None
    try:
        from PIL import ImageFont
        fn = ImageFont.load_default(28)
    except (TypeError, OSError):
        pass
    for i, c in enumerate(crops):
        d.text((8, 10 + 60 * i + 6), f"[{i + 1}]", fill=(0, 70, 200), font=fn)
        sheet.paste(c, (90, 10 + 60 * i))
        d.line([(0, 60 * i + 5), (W, 60 * i + 5)], fill=(200, 200, 200))
    return _png(sheet)


def vlm_read_blobs(a, boxes, model=None):
    if not boxes or len(boxes) > 30:
        return None
    j = parse_json(chat_vision(BLOB_PROMPT, [blob_sheet(a, boxes)], model))
    lab = j.get("labels") if isinstance(j.get("labels"), dict) else j
    out = []
    for i in range(len(boxes)):
        v = lab.get(str(i + 1))
        try:
            out.append(None if v is None else num(v))
        except (ValueError, TypeError):
            out.append(None)
    return out


def image_meta(iid):
    """자른 그림이면 {"parent": 원본 id, "box": [x0, y0, x1, y1]} (원본 좌표)"""
    p = os.path.join(WS, "images", iid + ".json")
    try:
        return json.loads(read(p)) if os.path.exists(p) else {}
    except (ValueError, OSError):
        return {}


def crop_image(iid, box, plot=None):
    """그림 일부를 새 그림으로 저장하고 원본·위치를 기록 (범례가 그래프 밖에 있을 때 원본에서 찾는다)"""
    a = image(iid)
    H, W = a.shape[:2]
    x0, y0, x1, y1 = [int(round(float(v))) for v in box]
    x0, x1 = sorted((max(0, min(W, x0)), max(0, min(W, x1))))
    y0, y1 = sorted((max(0, min(H, y0)), max(0, min(H, y1))))
    if x1 - x0 < 30 or y1 - y0 < 30:
        raise ValueError("고른 영역이 너무 작습니다")
    info = save_image(_png(Image.fromarray(a[y0:y1, x0:x1])))
    if info["id"] != iid:
        pm = image_meta(iid)  # 자른 그림을 또 자르면 맨 처음 원본 기준으로
        if pm.get("parent"):
            px0, py0 = pm["box"][:2]
            meta = {"parent": pm["parent"], "box": [px0 + x0, py0 + y0, px0 + x1, py0 + y1]}
        else:
            meta = {"parent": iid, "box": [x0, y0, x1, y1]}
        if plot:  # 그래프 고르기에서 찾은 축 상자 (자른 그림 좌표)
            meta["plot"] = [float(plot[0]) - x0, float(plot[1]) - y0, float(plot[2]) - x0, float(plot[3]) - y0]
        with open(ws("images", info["id"] + ".json"), "w", encoding="utf-8") as f:
            json.dump(meta, f)
        info["parent"] = meta["parent"]
    return info


TEXT_PROMPT = """Each numbered row [1], [2], ... of this image shows ONE short text label cut from a chart legend.
Copy the text of each row exactly as printed (write subscripts/superscripts as plain text, e.g. CO2, keep * and +).
Return ONLY JSON: {"labels": {"1": "text", "2": "text", ...}} with one entry per row; null if unreadable."""

CAT_PROMPT = """This is a bar chart. Every bar carries a magenta tag with a number [1], [2], ... drawn at its end and again at
its foot, just above the category axis.
For each tag give
 - "category": the category label printed on the category axis for that bar (under the x-axis for vertical bars,
   left of the y-axis for horizontal bars). Join a label printed on two lines with a space. If one label is centered
   under several bars, give it to each of those bars. Labels may be printed slanted (rotated): a slanted label belongs to
   the bar whose center is directly above the label's upper end (the end closest to the axis), so count bars and labels
   from left to right and keep them in the same order.
 - "group": a header shared by a group of bars (for example text printed inside the plot above bars that are
   separated by dashed lines), or "" if there is none.
 - "note": a short name printed right above or on that single bar (for example "Py*"), not a number or percentage; "" if none.
Return ONLY JSON: {"bars": {"1": {"category": "...", "group": "...", "note": "..."}, "2": {...}, ...}}"""


def vlm_read_texts(a, boxes, model=None):
    if not boxes or len(boxes) > 30:
        return None
    j = parse_json(chat_vision(TEXT_PROMPT, [blob_sheet(a, boxes)], model))
    lab = j.get("labels") if isinstance(j.get("labels"), dict) else j
    return [None if lab.get(str(i + 1)) in (None, "") else str(lab.get(str(i + 1))).strip() for i in range(len(boxes))]


def tagged_chart(a, bars, orient):
    """막대마다 번호 표를 붙인 그림 (VLM 이 범주 이름을 막대에 맞춰 읽도록)"""
    from PIL import ImageDraw, ImageFont
    im = Image.fromarray(a).convert("RGB")
    d = ImageDraw.Draw(im)
    fs = max(12, int(0.03 * min(im.size)))
    try:
        fn = ImageFont.load_default(fs)
    except (TypeError, OSError):
        fn = None
    for i, b in enumerate(bars):
        t = f"{i + 1}"
        tw, th = (fs * 0.6 * len(t) + 6, fs + 4)
        end = b["ends"][-1]
        if orient == "v":  # 막대 끝 위 + 바닥(범주 글자 바로 위) 두 곳
            spots = [(b["c"] - tw / 2, max(0, end - th - 4)), (b["c"] - tw / 2, b["base"] - th - 3)]
        else:
            spots = [(min(im.width - tw, end + 4), b["c"] - th / 2), (b["base"] + 3, b["c"] - th / 2)]
        for x, y in spots:
            d.rectangle([x, y, x + tw, y + th], fill=(220, 0, 200))
            d.text((x + 3, y + 1), t, fill=(255, 255, 255), font=fn)
    if max(im.size) > 1400:
        im.thumbnail((1400, 1400))
    return _png(im)


def _unique_cats(bars):
    """같은 범주 이름이 여럿이면(그룹마다 Conv.·Cir. 반복, 여러 막대에 걸친 이름) 막대 이름표 → 그룹 → 번호 순으로 구분"""
    def dup():
        c = {}
        for b in bars:
            c[b["cat"]] = c.get(b["cat"], 0) + 1
        return c
    for key in ("note", "group"):
        c = dup()
        for b in bars:
            extra = b.get(key) or ""
            if c[b["cat"]] > 1 and extra and extra not in b["cat"]:
                b["cat"] = f"{b['cat']} {extra}" if key == "note" else f"{extra} {b['cat']}"
    c, k = dup(), {}
    for b in bars:
        if c[b["cat"]] > 1:
            k[b["cat"]] = k.get(b["cat"], 0) + 1
            b["cat"] = f"{b['cat']} ({k[b['cat']]})"


def bar_auto(a, det, v=None, each=None, xvals=None, yvals=None, model=None, meta=None, orient=None, use_vlm=True):
    """막대그래프: 값축만 보정(눈금 값 짝짓기) + 막대·쌓인 조각·오차 막대 검출 + 범례 견본 짝짓기 + (VLM) 범주·계열 이름.
    → {"ok", "orient", "calib", "fit", "bars", "series", "warn"} 또는 {"ok": False, "error"}. 막대 위치는 픽셀(원본 좌표)."""
    ax, tk, lb = det["axes"], det["ticks"], det["labels"]
    H, W = a.shape[:2]
    orient, _ = B.find_bars(a, ax, orient if orient in ("v", "h") else None)
    k = "y" if orient == "v" else "x"
    each = each or {}
    vals = (yvals if k == "y" else xvals) or (v or {}).get(k + "_ticks") or []
    bx = ax["box"]
    shape = D.ticks_look_log(tk[k])
    hint = shape if shape is not None else (v or {}).get(k + "_log")
    bxs = lb.get(k + "_boxes") or []
    th = float(np.median([b[3] - b[1] for b in bxs])) if bxs else None
    args = dict(ticks=tk[k], majors=tk[k + "_major"], edges=(bx["x0"], bx["x1"]) if k == "x" else (bx["y0"], bx["y1"]),
                log=None, reverse=(k == "y"), hint=hint, text_h=th)
    opts = []
    if vals:
        c = D.pair_axis(vals, lb[k] or tk[k + "_major"], **args)
        if c:
            c["source"] = "list"
            opts.append(c)
    if each.get(k) and lb[k] and len(each[k]) == len(lb[k]):
        c = D.pair_axis(each[k], lb[k], fixed=True, **args)
        if c:
            c["source"] = "each"
            opts.append(c)
    mj = tk[k + "_major"]
    if vals and len(mj) == len(vals) >= 3:
        c = D.pair_axis(vals, mj, fixed=True, **args)
        if c:
            c["source"] = "ticks"
            opts.append(c)
    name = "y" if orient == "v" else "x"
    if not opts:
        return {"ok": False, "orient": orient,
                "error": f"막대그래프로 보이지만 값축({name}축) 눈금 숫자를 읽지 못했습니다 — 'AI 없이' 칸에 {name} 눈금 값을 넣거나 수동 보정하세요"}
    fit = min(opts, key=lambda c: (c["resid_px"] >= 1.5, c["ambiguous"], -c["n"], c["resid_px"]))
    cv = {"p1": fit["p1"], "p2": fit["p2"], "v1": fit["v1"], "v2": fit["v2"], "log": fit["log"]}
    an = B.analyze(a, ax, cv, orient)
    warn = []
    if fit["resid_px"] > 1.5:
        warn.append(f"{name}축 눈금 값과 위치가 잘 안 맞습니다 (잔차 {fit['resid_px']:.1f}px) — 확인하세요")
    if fit["dropped"]:
        warn.append(f"{name}축: 안 맞는 눈금 {', '.join(f'{x:g}' for x, _ in fit['dropped'])} 은(는) 빼고 맞췄습니다")
    xa, ya = ax["xaxis"], ax["yaxis"]
    if orient == "v":
        calib = {"x": {"p1": [bx["x0"], xa["y"]], "p2": [bx["x1"], xa["y"]], "v1": bx["x0"], "v2": bx["x1"], "log": False, "cat": True},
                 "y": {"p1": [ya["x"], fit["p1"]], "p2": [ya["x"], fit["p2"]], "v1": fit["v1"], "v2": fit["v2"], "log": fit["log"]}}
        base = xa["y"]
        conv = lambda p: p  # noqa: E731
    else:
        calib = {"x": {"p1": [fit["p1"], xa["y"]], "p2": [fit["p2"], xa["y"]], "v1": fit["v1"], "v2": fit["v2"], "log": fit["log"]},
                 "y": {"p1": [ya["x"], bx["y1"]], "p2": [ya["x"], bx["y0"]], "v1": bx["y1"], "v2": bx["y0"], "log": False, "cat": True}}
        base = ya["x"]
        conv = lambda p: W - p  # noqa: E731  돌린 좌표 → 원래 x
    if not an["bars"]:
        return {"ok": False, "orient": orient, "calib": calib, "fit": {name: fit},
                "error": "값축은 맞췄지만 막대를 찾지 못했습니다 — 그래프 영역을 다시 고르거나 '선·점' 방식으로 바꿔 보세요"}
    # 범례 견본 (그래프 안 → 없으면 자르기 전 원본에서 그래프 둘레)
    rects = []
    for b in an["bars"]:
        lo, hi = sorted((conv(b["top_px"]), base))
        rects.append([b["l"], lo, b["r"], hi] if orient == "v" else [lo, b["l"], hi, b["r"]])
    pats = B.legend_patches(a, avoid=rects)
    for p in pats:
        p["src"] = "image"
    pa = None
    if meta and meta.get("parent"):  # 자른 그림: 범례가 잘렸거나 밖에 있을 수 있음 → 원본 둘레에서도 (그래프 안 범례가 우선)
        try:
            pa = image(meta["parent"])
            x0, y0, x1, y1 = meta["box"]
            w_, h_ = x1 - x0, y1 - y0
            others = [q["box"] for q in B.find_panels(pa) if q["kind"] == "plot"]  # 다른 그래프 안(막대 조각이 견본처럼 보임)은 빼고
            pp = B.legend_patches(pa, avoid=[[r[0] + x0, r[1] + y0, r[2] + x0, r[3] + y0] for r in rects] + others,
                                  region=[x0 - 1.2 * w_, 0, x1 + 1.2 * w_, y1 + 0.2 * h_])  # 공유 범례는 흔히 그림 맨 위
            inside = [p for p in pp if x0 <= p["box"][0] and p["box"][2] <= x1 and y0 <= p["box"][1] and p["box"][3] <= y1]
            for p in pp:
                if p not in inside:
                    p.update(src="parent", penalty=15)
                    pats.append(p)
        except (OSError, ValueError):
            pa = None
    series = B.assign_series(an["bars"], pats)
    names = {}
    used = sorted({s["legend"] for s in series if s["legend"] is not None})
    if use_vlm and used:
        for srcname, img in (("image", a), ("parent", pa)):
            if img is None or not any(pats[j]["src"] == srcname for j in used):
                continue
            js = [j for j in range(len(pats)) if pats[j]["src"] == srcname][:30]  # 범례 전체를 한 장에 (일부만 보이면 잘 못 읽음)
            try:
                got = vlm_read_texts(img, [pats[j]["text"] for j in js], model) or []
                names.update({j: t.split("\n")[0].strip() for j, t in zip(js, got) if t})
            except Exception as e:  # 이름만 못 읽음 — 값은 그대로
                warn.append(f"범례 이름을 읽지 못했습니다 ({e}) — 계열 이름을 직접 고치세요")
    src = "parent" if any(pats[j]["src"] == "parent" for j in used) else ("image" if used else None)
    out_series = []
    for i, s in enumerate(series):
        j = s["legend"]
        nm = names.get(j) or f"계열 {i + 1}"
        out_series.append({"name": nm, "color": s["sig"]["color"], "kind": s["sig"]["kind"], "ink": s["sig"]["ink"],
                           "legend": j is not None, "legend_box": pats[j]["box"] if j is not None and pats[j]["src"] == "image" else None})
    bars = []
    for i, b in enumerate(an["bars"]):
        e = b["err"]
        bars.append({"cat": f"막대 {i + 1}", "group": "", "c": b["cx"], "w": [b["l"], b["r"]], "base": base,
                     "ends": [conv(sg["y1"]) for sg in b["segs"]], "series": [sg["series"] for sg in b["segs"]],
                     "err": [conv(e["hi_px"]), conv(e["lo_px"]) if e.get("lo_px") is not None else None] if e.get("hi_px") is not None else None,
                     "open": b["open"]})
    if use_vlm:
        try:
            j = parse_json(chat_vision(CAT_PROMPT, [tagged_chart(a, bars, orient)], model))
            got = j.get("bars") if isinstance(j.get("bars"), dict) else {}
            for i, b in enumerate(bars):
                it = got.get(str(i + 1)) or {}
                if isinstance(it, str):
                    it = {"category": it}
                if (it.get("category") or "").strip():
                    b["cat"] = str(it["category"]).strip()
                b["group"] = str(it.get("group") or "").strip()
                b["note"] = str(it.get("note") or "").strip()
        except Exception as e:
            warn.append(f"범주 이름을 읽지 못했습니다 ({e}) — 표에서 직접 고치세요")
    if use_vlm and orient == "v":  # 비스듬한 범주 글자: 막대 밑 위치로 짝짓고 하나씩 읽기 (전체 그림에서 읽으면 한 칸씩 밀리기 쉬움)
        try:
            sl = B.slanted_labels(a, ax, bars)
            if sl:
                got = vlm_read_texts(a, [x for x in sl if x], model) or []
                it = iter(got)
                for b, x in zip(bars, sl):
                    t = next(it, None) if x else None
                    if t:
                        b["cat"] = t.replace("\n", " ").strip()
        except Exception as e:
            warn.append(f"비스듬한 범주 글자를 읽지 못했습니다 ({e})")
    groups_of = {}
    for b in bars:
        groups_of.setdefault(b["cat"], set()).add(b.get("group") or "")
    for b in bars:  # 표에서 같은 범주 아래 나란한 막대(묶음 막대)를 한 줄로 모으는 열쇠 — 번호 붙이기 전 이름. 그룹 이름은 같은 범주가 여러 그룹에 있을 때만
        b["label"] = f"{b['group']} {b['cat']}".strip() if b.get("group") and len(groups_of[b["cat"]]) > 1 else b["cat"]
        b["key"] = b["label"] if not b["cat"].startswith("막대 ") else None
        if b["key"] is None:
            b.pop("key")
    _unique_cats(bars)
    nerr = sum(1 for b in bars if b["err"])
    return {"ok": True, "orient": orient, "calib": calib, "fit": {name: fit}, "bars": bars, "series": out_series, "warn": warn,
            "legend_src": src, "n_err": nerr}


def auto(iid, use_vlm=True, model=None, xvals=None, yvals=None, chart="auto", whole=False):
    """AI 자동 보정: 축·눈금·글자 위치 검출(결정론) + VLM 눈금 값 → 짝짓기·최소제곱. 계열 색 후보·범례 이름 제안.
    VLM 은 (1) 전체 그림에서 눈금 값 목록, (2) 글자 덩어리를 하나씩 번호 붙인 그림에서 덩어리별 값(1:1)을 읽고,
    둘 중 위치와 더 잘 맞는(잔차 작은) 쪽을 쓴다.
    chart: auto | xy(선·점) | bar(막대). 막대그래프면 값축만 보정하고 막대·조각·오차를 픽셀로 잰다(bar_auto).
    그래프가 여럿인 그림(논문 쪽)은 need_panel=True 와 panels 를 돌려준다 — 하나를 골라 자른 뒤 다시 부른다.
    늘 message(한국어: 무엇을 했고 다음에 무엇을 하면 되는지)를 채운다."""
    a = image(iid)
    H, W = a.shape[:2]
    meta = image_meta(iid)
    res = {"detect": None, "vlm": None, "error": None, "mode": "xy", "bar": None, "panels": [], "need_panel": False, "message": ""}
    panels = B.find_panels(a)
    res["panels"] = panels
    plots = [p for p in panels if p["kind"] == "plot"]
    big = [p for p in plots if (p["crop"][2] - p["crop"][0]) * (p["crop"][3] - p["crop"][1]) >= 0.5 * W * H]
    if not whole and (len(plots) >= 2 or (plots and not big)):
        res["need_panel"] = True
        skip = len(panels) - len(plots)
        res["message"] = (f"이 그림에서 그래프 {len(plots)}개를 찾았습니다" + (f" (지도·사진 등 {skip}개는 뺌)" if skip else "") +
                          " — 아래 목록에서 하나를 고르거나 ✂ 자르기로 영역을 끌어 고르세요.")
        res["detect"] = {"ok": False, "error": res["message"]}
        return res
    hint = dict(zip(("x0", "y0", "x1", "y1"), meta["plot"])) if meta.get("plot") else None
    det = D.auto_calibrate(a, None, None, box_hint=hint)
    res["detect"] = det
    if not det.get("ok"):
        res["message"] = "축선을 찾지 못했습니다 — ✂ 자르기로 그래프 영역만 고르거나, 수동 보정(축 위 두 점씩)을 쓰세요."
        det["error"] = res["message"]
        return res
    nb = len(B.find_bars(a, det["axes"])[1])
    det["bar_guess"] = nb
    brk = B.axis_break(a, det["axes"])
    xlog = ylog = None
    legend = []
    each = {}
    v = None
    if use_vlm and not (xvals and yvals):
        try:
            v = vlm_read(a, det, model)
            res["vlm"] = v
            xlog, ylog = v["x_log"], v["y_log"]
            legend = v["legend"]
            lb = det["labels"]
            for k in ("x", "y"):  # 글자 덩어리를 하나씩 다시 읽힘 → 개수 어긋남·모서리 겹침에 강함
                boxes = lb[k + "_boxes"] if k == "x" else lb["y_boxes"][::-1]  # y 는 아래 → 위
                if 2 <= len(boxes) <= 30:
                    try:
                        per = vlm_read_blobs(a, boxes, model)
                    except Exception:
                        per = None
                    if per and sum(p is not None for p in per) >= 2:
                        each[k] = per
                        v[k + "_ticks_each"] = per
        except Exception as e:
            res["error"] = f"VLM 읽기 실패 — 눈금 값을 직접 넣어 주세요 ({e})"
    vt = (v or {}).get("chart_type", "")
    xv = xvals or (v or {}).get("x_ticks") or []
    yv = yvals or (v or {}).get("y_ticks") or []
    want_bar = chart in ("bar", "vbar", "hbar") or (chart == "auto" and vt in ("bar", "hbar") and nb > 0)
    why = ""
    if not want_bar and (xv or each.get("x")) and (yv or each.get("y")):
        full = D.auto_calibrate(a, xv, yv, None, None, each, {"x": xlog, "y": ylog}, box_hint=hint)
        # VLM 이 말한 로그 여부와 잔차로 고른 결과가 다르면 경고 (결정은 잔차·눈금 모양)
        if full.get("calib"):
            for k, said in (("x", xlog), ("y", ylog)):
                got = full["calib"][k]["log"]
                if said is not None and said != got:
                    full["warn"].append(f"{k}축: VLM 은 {'로그' if said else '선형'}축이라 했지만 눈금 간격은 {'로그' if got else '선형'}축에 맞습니다 — 확인하세요")
        full["bar_guess"] = nb
        det = full
    if chart == "auto" and not want_bar and not det.get("calib") and nb > 0 and (v or xvals or yvals) and vt not in ("line", "scatter"):
        want_bar = True
        why = "x축(또는 y축)에 숫자 눈금이 없어(범주형) 막대 모드로 전환했습니다. "
    elif want_bar and chart == "auto":
        why = "막대그래프로 보여 막대 모드로 처리했습니다. "
    if want_bar:
        orient = {"vbar": "v", "hbar": "h"}.get(chart) or ("h" if vt == "hbar" else "v" if vt == "bar" else None)
        br = bar_auto(a, det, v, each, xvals, yvals, model, meta, orient, use_vlm=bool(v))
        res["bar"] = br
        if br.get("ok"):
            res["mode"] = "bar"
            det["calib"] = br["calib"]
            det["fit"] = None
            det["warn"] = list(det.get("warn") or []) + br["warn"]
            nseg = sum(len(b["ends"]) for b in br["bars"])
            res["message"] = (why + f"막대 {len(br['bars'])}개 ({'세로' if br['orient'] == 'v' else '가로'}), 조각 {nseg}개, "
                              f"계열 {len(br['series'])}개, 오차 막대 {br['n_err']}개를 찾았습니다. 결과 탭의 표를 확인하세요.")
        else:
            res["mode"] = "bar"
            res["message"] = why + br["error"]
    if not res["message"]:
        if det.get("calib"):
            res["message"] = "축 보정을 마쳤습니다 — 화면의 X1·X2·Y1·Y2 위치를 확인하고 데이터 탭에서 추출하세요."
        elif res["error"]:
            res["message"] = res["error"]
        elif not (use_vlm or xvals or yvals):
            res["message"] = (f"축·눈금을 찾았습니다 (x 눈금 {len(det['ticks']['x'])}개, y 눈금 {len(det['ticks']['y'])}개"
                              + (f", 막대 {nb}개" if nb else "") + "). ✨ AI 자동 보정을 누르거나 눈금 값을 넣으세요.")
        else:
            miss = [k for k, vv in (("x", xv or each.get("x")), ("y", yv or each.get("y"))) if not vv]
            res["message"] = ((f"{'·'.join(miss)}축 눈금 숫자를 읽지 못했습니다" if miss else "읽은 눈금 값과 눈금 위치를 짝짓지 못했습니다")
                              + " — 'AI 없이' 칸에 눈금 값을 직접 넣거나 수동 보정을 쓰세요. 막대그래프라면 방식을 '막대'로 바꿔 다시 누르세요.")
    if brk:
        det.setdefault("warn", []).insert(0, "y축이 끊겨 있습니다(위쪽 바깥에 끊김 표시와 다른 눈금 구간). 그림 영역(끊김 아래) 눈금으로만 보정했습니다 — "
                                         "끊김 위쪽 값은 ✂ 자르기로 그 구간만 골라 따로 보정하세요.")
        res["message"] += " ⚠ y축이 끊긴 그래프입니다 — 경고를 확인하세요."
        res["axis_break"] = brk
    det["colors"] = D.suggest_colors(a, det["box"])
    det["legend_boxes"] = D.find_legend_boxes(a, det["box"])
    det["series"] = D.match_legend(legend, det["colors"])
    if res["mode"] == "xy" and det.get("calib") and v and len(v.get("y2_ticks") or []) >= 2:
        try:
            det["calib2"] = right_axis(a, det, v, model)
        except Exception as e:
            det.setdefault("warn", []).append(f"오른쪽 y축 눈금을 맞추지 못했습니다 ({e})")
        if det.get("calib2"):
            c2 = det["calib2"]["y"]
            det.setdefault("warn", []).append(f"오른쪽 y축({c2['v1']:g}…{c2['v2']:g})이 따로 있습니다 — 범례에서 오른쪽 축이라고 읽은 계열은 오른쪽 눈금으로 바꿨습니다. "
                                              "다른 계열도 데이터 탭에서 '오른쪽 y축'을 켜고 끌 수 있습니다.")
    if res["mode"] == "xy" and det.get("calib") and (use_vlm or xvals or yvals):
        det["series"] = extract_all(a, det, chart_type=vt)
        got = [s for s in det["series"] if s.get("px")]
        if got:
            res["message"] = (res["message"].replace("데이터 탭에서 추출하세요", "결과 탭의 표를 확인하세요").replace("확인하고 결과", "확인하고, 결과") +
                              f" 계열 {len(got)}개를 자동으로 뽑았습니다 (" + ", ".join(f"{s['name']} {len(s['px'])}점" for s in got) + ").")
    res["detect"] = det
    return res


def right_axis(a, det, v, model=None):
    """둘째(오른쪽) y축 보정: 오른쪽 테두리 바깥의 눈금·글자 덩어리 + VLM 이 읽은 y2 눈금 값 → 왼쪽과 같은 x, 다른 y 를 쓰는 보정"""
    ax, calib = det["axes"], det["calib"]
    box = ax["box"]
    rl = D.right_labels(a, ax)
    if len(rl["y"]) < 2:
        return None
    dark = D.gray(a) < 150
    xr = int(round(box["x1"]))
    r0, r1 = int(round(box["y0"])), int(round(box["y1"]))
    L = max(12, int(0.03 * max(a.shape[:2])))
    out_ = dark[r0:r1 + 1, xr + 1:xr + 1 + L].T
    tk = [r0 + p for p, n in D._ticks_1d(D._run_len(out_), L - 1)] if out_.size else []
    vals = v.get("y2_ticks") or []
    opts = []
    args = dict(ticks=tk, majors=tk, edges=(box["y0"], box["y1"]), log=None, reverse=True, hint=None)
    c = D.pair_axis(vals, rl["y"], **args)
    if c:
        opts.append(c)
    boxes = rl["boxes"][::-1]
    if 2 <= len(boxes) <= 30:
        try:
            per = vlm_read_blobs(a, boxes, model)
        except Exception:
            per = None
        if per and sum(p is not None for p in per) >= 2:
            c = D.pair_axis(per, rl["y"], fixed=True, **args)
            if c:
                opts.append(c)
    if not opts:
        return None
    f = min(opts, key=lambda c: (c["resid_px"] >= 1.5, c["ambiguous"], -c["n"], c["resid_px"]))
    return {"x": calib["x"], "y": {"p1": [box["x1"], f["p1"]], "p2": [box["x1"], f["p2"]], "v1": f["v1"], "v2": f["v2"], "log": f["log"]}}


def assign_shapes(mg, sers):
    """표식 무리 ↔ 같은 색 범례 계열 짝: (1) VLM 이 읽은 표식 모양, (2) 범례 견본의 위→아래 순서, (3) 남은 것은 차례대로.
    → sers[i] 에 줄 무리 번호 목록"""
    k = len(sers)
    out = [None] * k
    used = set()
    for i, s in enumerate(sers):  # 모양 이름
        want = str(s.get("marker") or "").lower()
        for key in ("triangle", "square", "diamond", "circle"):
            if key in want:
                g = next((g for g in range(k) if g not in used and key in mg["shape"][g]), None)
                if g is not None:
                    out[i] = g
                    used.add(g)
                break
    leg = [g for g, _ in mg["legend"] if g not in used]
    if len(mg["legend"]) == k and len(set(g for g, _ in mg["legend"])) == k and not used:  # 견본이 다 보이면 위→아래 = 범례 순서
        return [g for g, _ in mg["legend"]]
    rest = [g for g in leg] + [g for g in range(k) if g not in used and g not in leg]
    for i in range(k):
        if out[i] is None:
            out[i] = rest.pop(0)
    return out


def extract_all(a, det, chart_type=""):
    """보정이 끝나면 제안 계열(범례 이름+색, 범례가 없으면 그림에서 찾은 색)을 모두 자동 추출 — 따로 ▶ 추출을 누르지 않아도 표가 찬다.
    범례가 없을 때 찾은 색 가운데 선·표식이 아닌 것(글자·격자 조각)은 버린다."""
    calib, box = det["calib"], det["box"]
    legends = det.get("legend_boxes")
    out = []
    named = any(s.get("name") and not s["name"].startswith("계열 ") for s in det.get("series") or [])
    W = box["x1"] - box["x0"]
    sers = [s for s in det.get("series") or []]
    # 같은 색 범례 계열이 여럿(흑백 그림): 표식 모양으로 가른다
    done = set()
    rgb = lambda c: np.array(D.hex2rgb(c), float)  # noqa: E731
    for i, s in enumerate(sers):
        if i in done or not s.get("color"):
            continue
        grp = [j for j, t in enumerate(sers) if t.get("color") and j not in done and np.sqrt(((rgb(t["color"]) - rgb(s["color"])) ** 2).sum()) < 60]
        if len(grp) < 2:
            continue
        done.update(grp)
        mg = D.marker_groups(a, calib, box, s["color"], len(grp), legends=legends)
        names = [sers[j]["name"] for j in grp]
        if not mg:
            for j in grp:
                sers[j]["warn"] = ["같은 색 계열이 여럿인데 표식 모양으로 가르지 못했습니다 — 데이터 탭에서 ⊘ 제외 영역·✎ 점 편집으로 나누세요"]
            continue
        order = assign_shapes(mg, [sers[j] for j in grp])
        xs_all = sorted(p[0] for g in mg["groups"] for p in g)
        grid = []
        for x in xs_all:  # 모든 무리의 x 를 모아 공통 x 격자 (같은 x 에서 쟀을 때)
            if grid and x - grid[-1][-1] <= 3:
                grid[-1].append(x)
            else:
                grid.append([x])
        grid = [float(np.mean(c)) for c in grid if len(c) >= 2]
        filled = {}
        for gi, g in enumerate(mg["groups"]):  # 다른 표식과 겹쳐 안 보인 점: 같은 x 격자에서 이웃 사이 직선으로 채움
            if len(g) < 3:
                continue
            gx = [p[0] for p in g]
            add = []
            for x in grid:
                if gx[0] < x < gx[-1] and min(abs(x - q) for q in gx) > 3:
                    add.append([x, float(np.interp(x, gx, [p[1] for p in g]))])
            if add:
                mg["groups"][gi] = sorted(g + add)
                filled[gi] = len(add)
        for gi, j in zip(order, grp):
            pts = mg["groups"][gi]
            sers[j] = dict(sers[j], px=pts, mode="line", how="markers", kind="line+markers", split=True,
                           warn=[f"같은 색 계열 {len(grp)}개({', '.join(names)})를 표식 모양으로 나눴습니다 — 이 계열: {mg['shape'][gi]}. 이름이 바뀌었으면 고치세요"]
                           + ([f"다른 표식과 겹쳐 안 보인 점 {filled[gi]}개는 이웃 점 사이 직선으로 채웠습니다"] if filled.get(gi) else []))
    for s in sers:
        if s.get("split") or s.get("warn"):
            out.append(s)
            continue
        if not s.get("color"):
            out.append(s)
            continue
        hint = "markers" if chart_type == "scatter" else ("markers" if "marker" in (s.get("style") or "") and chart_type != "line" else None)
        r = D.auto_series(a, calib, box, s["color"], legends=legends, hint=hint)
        px = r["px"]
        if not named:  # 범례 없음: 선이 그림 폭의 30% 넘게 지나거나 표식이 3개 이상일 때만 계열로
            xs = [p[0] for p in px]
            if not px or ((max(xs) - min(xs) < 0.3 * W) and len(px) < 3):
                continue
        o = dict(s, px=px, mode=r["mode"], how=r["how"], kind=r["kind"], warn=r["warn"])
        if det.get("calib2") and str(s.get("axis") or "").lower().startswith("r"):
            o["calib"] = det["calib2"]  # 오른쪽 y축 계열 — 값은 오른쪽 눈금으로
        out.append(o)
    for i, s in enumerate(out):
        if not s.get("name"):
            s["name"] = f"계열 {i + 1}"
    return out


# ── 추출·내보내기 ─────────────────────────────────────────────────────────
def extract(iid, calib, box, series, excludes=()):
    a = image(iid)
    legends = D.find_legend_boxes(a, box)
    r = D.run_series(a, calib, box, series, excludes, legends)
    mask = "data:image/png;base64," + base64.b64encode(D.mask_png(r["mask"], r["off"], a.shape, series.get("color") or "#000000")).decode()
    return {"px": r["px"], "data": r["data"], "warn": r["warn"], "mask": mask, "legend_boxes": legends}


def series_data(calib, s):
    """계열 하나 → 데이터 점 (x 순). 점은 늘 픽셀로 들고 다니고 보정으로 바꾼다 — 화면에 보이는 점 = 내보내는 점.
    오른쪽 y축 계열은 자기 보정(s["calib"])을 쓴다."""
    calib = s.get("calib") or calib
    px = s.get("px") or []
    if not px:
        return []
    xs, ys = D.to_data(calib, [p[0] for p in px], [p[1] for p in px])
    return sorted([[float(f"{x:.12g}"), float(f"{y:.12g}")] for x, y in zip(xs, ys)])


def fmt_num(v):
    return f"{v:.6g}"


def table(calib, series):
    cols = [(s.get("name") or f"계열{i + 1}", series_data(calib, s)) for i, s in enumerate(series)]
    head = []
    for name, _ in cols:
        head += [f"{name} x", f"{name} y"]
    n = max([len(d) for _, d in cols] or [0])
    rows = []
    for i in range(n):
        r = []
        for _, d in cols:
            r += [d[i][0], d[i][1]] if i < len(d) else ["", ""]
        rows.append(r)
    return head, rows, cols


def export(calib, series, fmt="csv", meta=None, bars=None, orient="v"):
    if bars:
        head, rows, long_ = bar_table(calib, bars, series, orient)
        cols = [("막대 조각", long_)]
    else:
        head, rows, cols = table(calib, series)
    base = re.sub(r'[\\/:*?"<>|\s]+', "_", ((meta or {}).get("name") or "digitized").strip()) or "digitized"
    if fmt == "xlsx":
        try:
            import openpyxl
        except ImportError:
            fmt = "csv"
        else:
            wb = openpyxl.Workbook()
            ws_ = wb.active
            ws_.title = "데이터"
            ws_.append(head)
            for r in rows:
                ws_.append([None if v == "" else v for v in r])
            for name, d in cols:
                sh = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "계열")
                sh.append(["범주", "계열", "아래 끝", "위 끝", "조각 값"] if bars else ["x", "y"])
                for r in d:
                    sh.append(list(r))
            info = wb.create_sheet("보정")
            for k in ("x", "y"):
                c = calib[k]
                if c.get("cat"):
                    info.append([f"{k}축", "범주축 (막대 위치)"])
                    continue
                info.append([f"{k}축", "로그" if c.get("log") else "선형", "점1(px)", *c["p1"], "값1", c["v1"], "점2(px)", *c["p2"], "값2", c["v2"]])
            for k, v in (meta or {}).items():
                if isinstance(v, (str, int, float)):
                    info.append([k, v])
            b = io.BytesIO()
            wb.save(b)
            return b.getvalue(), base + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    s = io.StringIO()
    w = csv.writer(s)
    w.writerow(head)
    for r in rows:
        w.writerow([fmt_num(v) if isinstance(v, float) else v for v in r])
    return ("﻿" + s.getvalue()).encode("utf-8"), base + ".csv", "text/csv; charset=utf-8"


def bar_values(calib, bar, orient):
    """막대 하나(픽셀) → (조각 위 끝 값들, 오차 [위, 아래] 값). 화면처럼 보정으로 바꾼다."""
    ends = bar.get("ends") or []
    if orient == "h":
        vs = D.to_data(calib, ends, [bar["c"]] * len(ends))[0]
    else:
        vs = D.to_data(calib, [bar["c"]] * len(ends), ends)[1]
    err = None
    if bar.get("err"):
        hi, lo = bar["err"]
        top = float(vs[-1])
        p = [hi, lo if lo is not None else None]
        q = [x for x in p if x is not None]
        ev = (D.to_data(calib, q, [bar["c"]] * len(q))[0] if orient == "h" else D.to_data(calib, [bar["c"]] * len(q), q)[1]).tolist()
        h_ = ev[0]
        l_ = ev[1] if len(ev) > 1 else 2 * top - h_
        err = [h_, l_]
    return [float(x) for x in vs], err


def bar_table(calib, bars, series, orient="v"):
    """범주 × 계열 표 (화면 barRows 와 같은 규칙): 같은 범주 이름(그룹+이름) 아래 나란한 막대는 한 줄에 계열별로.
    쌓인 막대가 있으면 '막대 끝(합계)', 오차 막대가 있는 계열마다 '오차 +/−'. long: 조각마다 한 줄"""
    names = [s.get("name") or f"계열 {i + 1}" for i, s in enumerate(series)]
    nm = lambda j: names[j] if j < len(names) else f"계열 {j + 1}"  # noqa: E731
    used = sorted({j for b in bars for j in b.get("series") or []})
    stacked = any(len(b.get("ends") or []) > 1 for b in bars)
    err_s = sorted({b["series"][-1] for b in bars if b.get("err")})
    rows, idx, long_ = [], {}, []
    for b in bars:
        tops, err = bar_values(calib, b, orient)
        seg, prev = {}, 0.0
        for j, t in zip(b.get("series") or [], tops):
            seg[j] = seg.get(j, 0.0) + t - prev
            long_.append([b.get("cat") or "", nm(j), prev, t, t - prev])
            prev = t
        key = b.get("key", b.get("cat"))
        r = idx.get(key)
        if r is None or any(j in r["seg"] for j in seg):
            r = {"cat": b.get("label") if b.get("key") is not None else b.get("cat"), "seg": {}, "tops": [], "err": {}}
            if any(q["cat"] == r["cat"] for q in rows):
                r["cat"] = b.get("cat")
            rows.append(r)
            idx[key] = r
        r["seg"].update(seg)
        r["tops"].append(tops[-1] if tops else "")
        if err:
            r["err"][b["series"][-1]] = [err[0] - tops[-1], tops[-1] - err[1]]
    head = ["범주"] + [nm(j) for j in used] + (["막대 끝(합계)"] if stacked else []) + [f"{nm(j)} 오차 {c}" for j in err_s for c in "+−"]
    out = [[r["cat"]] + [r["seg"].get(j, "") for j in used] + ([r["tops"][0] if len(r["tops"]) == 1 else ""] if stacked else [])
           + [v for j in err_s for v in r["err"].get(j, ["", ""])] for r in rows]
    return head, out, long_


# ── 이력 (WORKSPACE/projects/<id>.json + history.jsonl) ───────────────────
def save_project(p):
    pid = p.get("pid") if re.fullmatch(r"\d{4}-\d{2}-\d{2}-[0-9a-f]{6}", p.get("pid") or "") else \
        f"{datetime.date.today()}-{secrets.token_hex(3)}"
    p = dict(p, pid=pid, saved=datetime.datetime.now().isoformat(timespec="seconds"))
    with LOCK:
        with open(ws("projects", pid + ".json"), "w", encoding="utf-8") as f:
            json.dump(p, f, ensure_ascii=False)
    return pid


def history(limit=200):
    d = os.path.join(WS, "projects")
    if not os.path.isdir(d):
        return []
    out = []
    for fn in sorted(os.listdir(d), key=lambda f: os.path.getmtime(os.path.join(d, f)), reverse=True)[:limit]:
        try:
            p = json.loads(read(os.path.join(d, fn)))
        except (ValueError, OSError):
            continue
        out.append({"pid": p["pid"], "name": p.get("name") or "", "saved": p.get("saved"), "image": p.get("image", {}).get("id"),
                    "series": [{"name": s.get("name"), "n": len(series_data(p["calib"], s)) if p.get("calib") else len(s.get("px") or [])}
                               for s in p.get("series") or []]})
    return out


def load_project(pid):
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}-[0-9a-f]{6}", pid or ""):
        raise ValueError("잘못된 id")
    return json.loads(read(os.path.join(WS, "projects", pid + ".json")))


def models():
    try:
        if LLM_API == "openai":
            req = urllib.request.Request(LLM_BASE + "/models", headers={"Authorization": f"Bearer {LLM_KEY}"} if LLM_KEY else {})
            with urllib.request.urlopen(req, timeout=3) as r:
                return [m["id"] for m in json.load(r)["data"]]
        with urllib.request.urlopen(LLM_BASE + "/api/tags", timeout=3) as r:
            return [m["name"] for m in json.load(r)["models"]]
    except Exception:
        return []


# ── HTTP ───────────────────────────────────────────────────────────────
HTML = read(os.path.join(ROOT, "ui.html")) if os.path.exists(os.path.join(ROOT, "ui.html")) else "ui.html 없음"

# ── 저작권 표기 (LICENSE·NOTICE 참고) ─────────────────────────────────────
_SIG = __import__("base64").b64decode("wqkgMjAyNiBnZ2dnODY1NyDCtyBkb25nanVraW0uZGV2QGdtYWlsLmNvbQ==").decode()
_SIG_A = __import__("base64").b64decode("Z2dnZzg2NTcgPGRvbmdqdWtpbS5kZXZAZ21haWwuY29tPg==").decode()


def signed(html):
    """화면에 저작권 표기를 붙인다. ui.html 에서 지워져도 서버가 내보낼 때 다시 붙는다."""
    name, mail = _SIG.split(" · ")
    if 'name="author"' not in html:
        meta = f'<meta name="author" content="{name[7:]} <{mail}>">'
        html = html.replace("<head>", "<head>" + meta, 1) if "<head>" in html else meta + html
    if "data-sig" not in html:
        tag = (f'<!-- {_SIG} --><div data-sig title="{mail}" style="text-align:center;font-size:11px;color:#9aa0a6;'
               f'opacity:.55;margin:6px 0 2px">{name}</div>')
        html = html.replace("</body>", tag + "</body>", 1) if "</body>" in html else html + tag
    return html


def data_url_bytes(s):
    return base64.b64decode((s or "").split(",", 1)[-1])


class H(BaseHTTPRequestHandler):
    def log_message(self, fmt, *a):
        if any(k in (a[0] if a else "") for k in ("/api/auto", "/api/export")):
            super().log_message(fmt, *a)

    def _send(self, body, ctype="application/json", code=200, filename=None):
        b = body if isinstance(body, bytes) else json.dumps(body, ensure_ascii=False, default=_json_default).encode()
        self.send_response(code)
        self.send_header("X-Author", _SIG_A)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(b)))
        self.send_header("Cache-Control", "no-store")
        if filename:
            self.send_header("Content-Disposition", "attachment; filename*=UTF-8''" + urllib.parse.quote(filename))
        self.end_headers()
        self.wfile.write(b)

    def do_GET(self):
        u = urllib.parse.urlparse(self.path)
        path, q = u.path, dict(urllib.parse.parse_qsl(u.query))
        try:
            if path == "/api/health":
                return self._send({"ok": True})
            if path == "/api/meta":
                return self._send({"model": MODEL, "llm": LLM_BASE, "models": models(),
                                   "pdf": bool(shutil.which("pdftoppm")), "xlsx": _has("openpyxl")})
            if path == "/api/history":
                return self._send(history())
            m = re.fullmatch(r"/api/project/([\w-]+)", path)
            if m:
                return self._send(load_project(m.group(1)))
            m = re.fullmatch(r"/api/img/([0-9a-f]{16})\.png", path)
            if m:
                with open(os.path.join(WS, "images", m.group(1) + ".png"), "rb") as f:
                    return self._send(f.read(), "image/png")
            m = re.fullmatch(r"/api/pdfpage/([0-9a-f]{16})/(\d+)\.png", path)
            if m:
                return self._send(pdf_render(m.group(1), m.group(2), q.get("dpi", 40)), "image/png")
            self._send(signed(HTML).encode(), "text/html; charset=utf-8")
        except FileNotFoundError as e:
            self._send({"error": str(e) or "없음"}, code=404)
        except Exception as e:
            self._send({"error": f"{type(e).__name__}: {e}"}, code=500)

    def do_POST(self):
        try:
            n = int(self.headers.get("Content-Length") or 0)
            if n > MAX_UPLOAD * 1.4:
                return self._send({"error": "파일이 너무 큽니다 (40MB 이하)"}, code=413)
            req = json.loads(self.rfile.read(n) or b"{}")
            if self.path == "/api/image":
                return self._send(save_image(data_url_bytes(req.get("data"))))
            if self.path == "/api/pdf":
                return self._send(save_pdf(data_url_bytes(req.get("data"))))
            if self.path == "/api/pdfpage":
                png = pdf_render(req.get("id"), req.get("page") or 1, req.get("dpi") or 200)
                return self._send(save_image(png))
            if self.path == "/api/auto":
                return self._send(auto(req.get("id"), req.get("vlm", True), req.get("model"),
                                       [num(v) for v in req.get("xvals") or []], [num(v) for v in req.get("yvals") or []],
                                       req.get("chart") or "auto", bool(req.get("whole"))))
            if self.path == "/api/crop":
                return self._send(crop_image(req.get("id"), req.get("box") or [], req.get("plot")))
            if self.path == "/api/panels":
                return self._send({"panels": B.find_panels(image(req.get("id")))})
            if self.path == "/api/colors":
                return self._send({"colors": D.suggest_colors(image(req.get("id")), req["box"])})
            if self.path == "/api/extract":
                return self._send(extract(req.get("id"), req["calib"], req["box"], req.get("series") or {}, req.get("excludes") or ()))
            if self.path == "/api/export":
                body, name, ctype = export(req["calib"], req.get("series") or [], req.get("format") or "csv", req.get("meta"),
                                           req.get("bars"), req.get("orient") or "v")
                return self._send(body, ctype, filename=name)
            if self.path == "/api/save":
                return self._send({"pid": save_project(req.get("project") or {})})
            if self.path == "/api/history/delete":
                pid = req.get("pid") or ""
                if re.fullmatch(r"\d{4}-\d{2}-\d{2}-[0-9a-f]{6}", pid):
                    p = os.path.join(WS, "projects", pid + ".json")
                    if os.path.exists(p):
                        os.remove(p)
                return self._send({"ok": True})
            self._send({"error": "없는 경로"}, code=404)
        except FileNotFoundError as e:
            self._send({"error": str(e)}, code=404)
        except Exception as e:
            self._send({"error": f"{type(e).__name__}: {e}"}, code=500)


def _has(mod):
    try:
        __import__(mod)
        return True
    except ImportError:
        return False


def _json_default(o):
    if isinstance(o, np.generic):
        return o.item()
    if isinstance(o, np.ndarray):
        return o.tolist()
    if isinstance(o, tuple):
        return list(o)
    raise TypeError(type(o).__name__)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--cli":
        a = D.load(sys.argv[2])
        calib = json.load(open(sys.argv[3], encoding="utf-8"))
        ss = json.load(open(sys.argv[4], encoding="utf-8"))
        ss = ss if isinstance(ss, list) else [ss]
        box = calib.get("box") or D.detect_axes(a)["box"]
        for s in ss:  # 화면과 같게: 점은 픽셀로 들고 내보낼 때 보정으로 바꾼다
            r = D.run_series(a, calib, box, s)
            if s.get("mode") == "points" or not r["data"]:
                s["px"] = r["px"]
            else:
                px, py = D.to_pixel(calib, [d[0] for d in r["data"]], [d[1] for d in r["data"]])
                s["px"] = list(zip(px.tolist(), py.tolist()))
        sys.stdout.write(export(calib, ss, "csv")[0].decode("utf-8-sig"))
        sys.exit(0)
    print(f"digitizer local → http://0.0.0.0:{PORT}  (vlm={LLM_API} {LLM_BASE} {MODEL}, workspace={WS})  {_SIG}")
    ThreadingHTTPServer(("0.0.0.0", PORT), H).serve_forever()
